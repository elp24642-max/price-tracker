import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HISTORY_FILE = "history.json"


def load_history(path=HISTORY_FILE):
    """Цены с прошлого запуска: {url: price}. Если файла нет, значит запуск первый."""
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as file:
        return json.load(file)


def save_history(books, path=HISTORY_FILE):
    with open(path, "w", encoding="utf-8") as file:
        json.dump({book["url"]: book["price"] for book in books}, file, ensure_ascii=False)


def find_changes(books, history):
    """Книги, у которых цена отличается от прошлого запуска."""
    changes = []
    for book in books:
        old = history.get(book["url"])
        if old is not None and old != book["price"]:
            changes.append({**book, "old_price": old, "diff": round(book["price"] - old, 2)})
    return changes


def _fill_sheet(sheet, headers, rows):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        width = max(len(str(cell.value or "")) for cell in sheet[get_column_letter(index)])
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 60)


def write_excel(books, changes, path):
    workbook = Workbook()
    prices = workbook.active
    prices.title = "Цены"
    _fill_sheet(
        prices,
        ["Название", "Цена", "В наличии", "Ссылка"],
        [[b["title"], b["price"], "да" if b["in_stock"] else "нет", b["url"]] for b in books],
    )
    changed = workbook.create_sheet("Изменения цен")
    _fill_sheet(
        changed,
        ["Название", "Было", "Стало", "Разница", "Ссылка"],
        [[c["title"], c["old_price"], c["price"], c["diff"], c["url"]] for c in changes],
    )
    workbook.save(path)
