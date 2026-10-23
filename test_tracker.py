import os
import tempfile
import unittest

from openpyxl import load_workbook

from report import find_changes, write_excel
from scraper import HEADERS, parse_page, parse_price

HTML = """
<article class="product_pod">
  <h3><a href="a-light_1000/index.html" title="A Light in the Attic">A Light in the ...</a></h3>
  <div class="product_price">
    <p class="price_color">£51.77</p>
    <p class="instock availability"><i class="icon-ok"></i> In stock </p>
  </div>
</article>
<article class="product_pod">
  <h3><a href="tipping_999/index.html" title="Tipping the Velvet">Tipping the ...</a></h3>
  <div class="product_price">
    <p class="price_color">£53.74</p>
    <p class="availability"> Out of stock </p>
  </div>
</article>
"""
PAGE_URL = "https://books.toscrape.com/catalogue/page-1.html"


class ParseTest(unittest.TestCase):
    def test_parse_price(self):
        self.assertEqual(parse_price("£51.77"), 51.77)
        self.assertEqual(parse_price("Â£10"), 10.0)  # битая кодировка не ломает разбор
        with self.assertRaises(ValueError):
            parse_price("бесплатно")

    def test_parse_page(self):
        books = parse_page(HTML, PAGE_URL)
        self.assertEqual(len(books), 2)
        self.assertEqual(books[0]["title"], "A Light in the Attic")  # полное название
        self.assertEqual(books[0]["price"], 51.77)
        self.assertTrue(books[0]["in_stock"])
        self.assertFalse(books[1]["in_stock"])
        self.assertEqual(
            books[0]["url"], "https://books.toscrape.com/catalogue/a-light_1000/index.html"
        )


class HeadersTest(unittest.TestCase):
    def test_headers_are_latin1(self):
        # Кириллица в заголовке роняет requests ещё до отправки запроса
        for value in HEADERS.values():
            value.encode("latin-1")


class ReportTest(unittest.TestCase):
    def setUp(self):
        self.books = parse_page(HTML, PAGE_URL)

    def test_first_run_has_no_changes(self):
        self.assertEqual(find_changes(self.books, {}), [])

    def test_detects_price_change(self):
        history = {self.books[0]["url"]: 49.99, self.books[1]["url"]: 53.74}
        changes = find_changes(self.books, history)
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]["old_price"], 49.99)
        self.assertEqual(changes[0]["diff"], 1.78)

    def test_excel_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "prices.xlsx")
            write_excel(self.books, [], path)
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["Цены", "Изменения цен"])
            self.assertEqual(workbook["Цены"].max_row, 3)  # заголовок + 2 книги
            self.assertEqual(workbook["Цены"]["B2"].value, 51.77)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
